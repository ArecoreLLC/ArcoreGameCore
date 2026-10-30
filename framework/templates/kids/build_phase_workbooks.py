"""Create three phase-specific worksheets per grade; initial design uses the original workbook.

Requires python-docx. Rebuilding overwrites generated files, so keep completed
worksheets in a separate location.
"""
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from build_workbooks import LEVELS, box, para, heading, flow

ROOT = Path(__file__).resolve().parent

# Each grade gets a different feature example and a concrete paper test.
FEATURES = [
 ('A new sound', 'Touch a bell', 'Bell ready?', 'Play a chime', 'Wait for the bell to be ready', 'Point to the bell. Make its sound when it is ready. Try when it is not ready.'),
 ('A helpful object', 'Pick up a lantern', 'Is it dark?', 'Light the path', 'Keep the lantern for later', 'Act out picking up the lantern in daylight and in darkness.'),
 ('A new move', 'Press the jump button', 'Standing on the ground?', 'Jump over the puddle', 'Wait until the hero lands', 'Use a paper hero. Try jumping from the ground and while already in the air.'),
 ('A clue finder', 'Use the magnifying glass', 'Near a hidden clue?', 'Reveal the clue', 'Say no clue is nearby', 'Try near a clue and far away. Check the clue is still visible after finding it.'),
 ('A reward with a cost', 'Buy a seed', 'Have at least 3 coins?', 'Spend 3 coins and give one seed', 'Show the missing coins', 'Try 2, 3, and 4 coins. Count coins and seeds after each purchase.'),
 ('A choice with a consequence', 'Choose a gift', 'Gift is for the keeper?', 'Increase keeper trust', 'Give the gift to the village', 'Try both gifts. Record who benefits and what changes in the next scene.'),
 ('A crafting recipe', 'Choose craft', 'Have wood and rope?', 'Spend the ingredients and make a raft', 'Name the missing ingredient', 'Try each ingredient alone, neither, and both. Check one craft spends one recipe.'),
 ('A changing companion', 'Ask companion to scout', 'Companion is rested?', 'Set state to scouting and reveal a clue', 'Show that the companion needs rest', 'Try rested, scouting, and tired states. Do not send a scouting companion twice.'),
 ('A choice remembered later', 'Return to the village', 'Bridge was repaired?', 'Show the trader arriving', 'Show the trader waiting across the river', 'Try two playthroughs with different bridge choices. Check the later scene remembers each.'),
 ('A resource tradeoff', 'Start a rescue', 'Supplies >= 2 and rescue not started?', 'Spend 2 supplies and start rescue once', 'Explain missing supplies or an active rescue', 'Try 1, 2, and 3 supplies, then request again. Check the cost is applied only once.'),
 ('A feature with a measurable goal', 'Request a hint', 'Puzzle unsolved and hint not shown?', 'Show one clue and mark hint shown', 'Keep current feedback without repeating the hint', 'Try unsolved, solved, and repeated requests. Observe whether a tester can continue after the hint.'),
]

FEATURE_CHALLENGES = [
 'Draw the new thing and point to what makes it happen.',
 'Show what the player does and what changes on the screen.',
 'Show where the new move helps and where it should not work.',
 'Explain how the player learns the new tool exists.',
 'Write the cost and reward. What stops the player getting a free reward?',
 'Connect the feature to a story choice with two different results.',
 'Show how the feature fits into explore, collect, build, and return.',
 'Name the states before and after use. Show how the feature becomes ready again.',
 'Track one saved choice and the later scene that reads it.',
 'Name the resources this feature reads and changes. Plan repeated-use behavior.',
 'Write an observable success goal and the smallest version that can test it.',
]

CHARACTER_CHALLENGES = [
 'Draw a happy face and a surprised face for your character.',
 'Draw how your character moves and one thing they carry.',
 'Draw something that makes your character easy to spot.',
 'Draw a pose that shows a feeling without using words.',
 'Show how one strength helps and one limit creates a challenge.',
 'Show a choice that reveals what the character values.',
 'Show how the character ability helps the main gameplay loop.',
 'Draw the character in two states. Label what changes between them.',
 'Show how the character reacts differently to two earlier player choices.',
 'Show the cost, benefit, and story consequence of the character ability.',
 'Show the silhouette, movement, and feedback that communicate this character role.',
]

CHARACTER_RULES = [
 ('Meet a friend', 'Friend looks sad?', 'Offer a hug', 'Wave hello'),
 ('See a dark path', 'Have a lantern?', 'Hold the lantern up', 'Ask a friend for light'),
 ('Try to glide', 'Wings are open?', 'Float down gently', 'Land and try again'),
 ('Hear a strange sound', 'Feeling brave?', 'Investigate the sound', 'Ask a friend to come along'),
 ('Lift a fallen branch', 'Energy at least 2?', 'Spend 2 energy and clear the path', 'Rest to regain energy'),
 ('Hear a request for help', 'Player chooses to help?', 'Join the rescue', 'Explain why the character leaves'),
 ('Use repair ability', 'Have a part?', 'Spend one part and repair a machine', 'Ask for a missing part'),
 ('Receive a warning', 'State is calm?', 'Set state to alert and look around', 'Keep current state and acknowledge warning'),
 ('Meet the player again', 'Player kept the promise?', 'Offer a secret route', 'Offer a chance to rebuild trust'),
 ('Use a powerful ability', 'Energy >= 3 and ability ready?', 'Spend 3 energy and start cooldown', 'Show missing energy or remaining cooldown'),
 ('Request companion help', 'Companion available and task valid?', 'Assign task and signal the new state', 'Explain why the task cannot start'),
]

def setup(grade,phase,i):
    d=Document(); s=d.sections[0]
    s.page_width=Inches(8.5); s.page_height=Inches(11)
    s.top_margin=s.bottom_margin=Inches(.6)
    s.left_margin=s.right_margin=Inches(.75)
    normal=d.styles['Normal']; normal.font.name='Arial'
    normal.font.size=Pt(13 if i<3 else 11)
    normal.paragraph_format.space_after=Pt(6)
    for name,size in [('Title',22),('Heading 1',17)]:
        style=d.styles[name]; style.font.name='Arial'; style.font.size=Pt(size)
        style.font.color.rgb=RGBColor(0,0,0)
    for style in d.styles:
        for border in list(style.element.iter(qn('w:pBdr'))): border.getparent().remove(border)
    d.core_properties.title=f'{grade} {phase}'
    d.core_properties.author='Arcore3dGameDesignFramework'
    d.core_properties.subject='Printable phase-specific game design worksheet'
    d.add_paragraph(f'{grade} {phase}','Title')
    p=s.footer.paragraphs[0]
    p.add_run(f'{grade}  |  {phase}  |  ')
    field=OxmlElement('w:fldSimple'); field.set(qn('w:instr'),'PAGE'); p._p.append(field)
    for run in p.runs: run.font.size=Pt(9)
    return d

def identity(d,i,opening):
    para(d,opening)
    if i<3: para(d,'Draw, point, or tell a helper. Your helper can write your words.',11)
    para(d,'Designer ____________________    Date ______________',11)
    para(d,'Game _______________________    Version ___________',11)

def rules_page(d,phase,rule,test,i):
    d.add_page_break(); heading(d,'How it works')
    para(d,'A trigger starts a rule. A condition asks a yes or no question. An action is what happens next.',11)
    para(d,'Example',11,True); flow(d,*rule)
    para(d,'Make your own rule. Follow the arrows and fill in both paths.',11)
    box(d,'WHEN → What starts this rule?',.42)
    box(d,'↓ CHECK → What yes or no question does the game ask?',.42)
    box(d,'YES → What happens?',.42)
    box(d,'NO → What happens instead?',.42)
    heading(d,'Try it with a tester')
    para(d,test,11)
    box(d,'My test and the result I expect',.42)
    box(d,'What actually happened when we tried it',.42)
    box(d,'One change to try next',.42)

def make_feature(i,data):
    grade=data[0]; focus,*rest=FEATURES[i]; trigger,condition,yes,no,test=rest
    d=setup(grade,'New Feature Design',i)
    identity(d,i,'Add one new thing to a game you already have. Start with a small version you can try on paper.')
    para(d,f'Idea to try: {focus.lower()}. You can choose a different feature.',11)
    if i<5:
        qs=['What new thing can the player do or see?', 'Why would this make the game more fun?', 'What must still work the same way?', 'What is the smallest version we can try first?']
    else:
        qs=['What new behavior are you adding and why does the player need it?', 'Where does it fit into the existing story or gameplay?', 'What existing behavior must keep working?', 'What will you build first and what can wait?']
    for q in qs: box(d,q,.48)
    box(d,'Draw before and after. Label what changes and what stays the same.',1.45)
    para(d,FEATURE_CHALLENGES[i],11)
    rules_page(d,'feature',(trigger,condition,yes,no),test,i)
    boxcount=12
    return d,boxcount

def make_level(i,data):
    grade=data[0]; d=setup(grade,'New Level Design',i)
    identity(d,i,'Plan a new place to play in your game. Give the player a clear start, a challenge, and a way to finish.')
    if i<5:
        qs=['What is this place called and why does the hero go there?', 'What does the player need to do to finish?', 'What gets in the way and what helps?', 'What happens if the player needs another try?']
    else:
        qs=['What story event brings the player here and what changes on leaving?', 'What is the completion goal and how will the player recognize it?', 'How will you introduce, practice, then challenge one skill?', 'Where can the player recover or restart after a mistake?']
    for q in qs: box(d,q,.48)
    box(d,'Draw the level. Mark START, GOAL, paths, obstacles, and helpful objects.',1.6)
    para(d,data[12]+' Add arrows to show the route.',11)
    test=data[13]+' Check the start, the finish, and a retry.'
    rules_page(d,'level',tuple(data[8:12]),test,i)
    return d,12

def make_character(i,data):
    grade=data[0]; d=setup(grade,'Character Design',i)
    identity(d,i,'Create a character who belongs in your game. Decide what they want, how they look, and what they do.')
    if i<5:
        qs=['What is the character name and are they a hero, friend, or troublemaker?', data[2], 'What does this character want and what gets in the way?', 'What special thing can they do and when do they need help?']
    else:
        qs=['What is the character name and their role in the story and gameplay?', data[2], 'What do they want and how does that conflict with another character?', 'What can they do and what limit keeps their ability fair?']
    for q in qs: box(d,q,.48)
    box(d,'Draw the character. Label important shapes, colors, objects, and movement.',1.6)
    para(d,CHARACTER_CHALLENGES[i],11)
    test=('Act as your character. Try the YES and NO paths. Ask your tester what the character wants.' if i<5 else
          'Act out the ability in an allowed and a blocked case. Repeat it. Ask a tester what the reactions tell them about the character.')
    rules_page(d,'character',CHARACTER_RULES[i],test,i)
    return d,12

PHASES=[('new-features',make_feature),('new-levels',make_level),('characters',make_character)]

def build():
    for folder,maker in PHASES:
        dest=ROOT/folder; dest.mkdir(exist_ok=True)
        title={'new-features':'New_Feature_Design','new-levels':'New_Level_Design','characters':'Character_Design'}[folder]
        for i,data in enumerate(LEVELS):
            prefix='Kindergarten' if i==0 else f'Grade_{i:02d}'
            d,expected=maker(i,data)
            assert len(d.tables)==expected, (folder,i,len(d.tables))
            path=dest/f'{prefix}_{title}.docx'; d.save(path)
            print(path.relative_to(ROOT))

if __name__=='__main__': build()
